import openpyxl
import re
from datetime import datetime, date
from database import get_db, normalizar_numero_celular


def limpiar_texto(txt):
    """
    Limpia y normaliza un texto: elimina espacios adicionales, 
    caracteres especiales de salto de línea y tildes, convirtiéndolo a mayúsculas.
    """
    if txt is None:
        return ""
    txt = str(txt).strip().upper()
    txt = txt.replace("Á", "A").replace("É", "E").replace("Í", "I").replace("Ó", "O").replace("Ú", "U")
    txt = re.sub(r'[\r\n\t]+', ' ', txt)
    txt = re.sub(r'\s+', ' ', txt)
    return txt.strip()


def normalizar_valor_excel(valor):
    """
    Normaliza valores leídos desde Excel para evitar valores flotantes que son enteros,
    y convertir nulos o 'N/A' a cadenas vacías.
    """
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    texto = str(valor).strip()
    if texto.upper() in ("NONE", "NULL", "N/A", "-", ""):
        return ""
    return texto.strip()


def formatear_solo_fecha(val):
    """
    Toma un valor que puede ser datetime o un string y lo convierte 
    a un formato estándar de fecha dd/mm/yyyy.
    """
    if val is None:
        return ""
    if isinstance(val, (datetime, date)):
        return val.strftime("%d/%m/%Y")

    val_str = str(val).strip()
    if val_str.upper() in ("NONE", "NULL", "N/A", "-") or not val_str:
        return ""

    # Intenta hacer match con formato ISO YYYY-MM-DD
    match_iso = re.match(r'^(\d{4})-(\d{1,2})-(\d{1,2})', val_str)
    if match_iso:
        anio, mes, dia = match_iso.group(1), match_iso.group(2).zfill(2), match_iso.group(3).zfill(2)
        return f"{dia}/{mes}/{anio}"

    # Intenta hacer match con formato latino DD/MM/YYYY
    match_latam = re.match(r'^(\d{1,2})/(\d{1,2})/(\d{4})', val_str)
    if match_latam:
        dia, mes, anio = match_latam.group(1).zfill(2), match_latam.group(2).zfill(2), match_latam.group(3)
        return f"{dia}/{mes}/{anio}"

    # Extraer dígitos para intentar inferir la fecha
    nums = re.findall(r'\d+', val_str)
    if len(nums) >= 3:
        if len(nums[0]) == 4:
            dia, mes, anio = nums[2], nums[1], nums[0]
            return f"{dia.zfill(2)}/{mes.zfill(2)}/{anio}"
        if len(nums) >= 3 and len(nums[-1]) == 4:
            dia, mes, anio = nums[0], nums[1], nums[-1]
            return f"{dia.zfill(2)}/{mes.zfill(2)}/{anio}"

    return val_str


def fecha_iso_importada(valor):
    if valor is None or str(valor).strip().upper() in ("", "NONE", "NULL", "N/A", "-"):
        return ""
    if isinstance(valor, (datetime, date)):
        return valor.date().isoformat() if isinstance(valor, datetime) else valor.isoformat()
    if isinstance(valor, (int, float)):
        try:
            convertido = openpyxl.utils.datetime.from_excel(valor)
        except (TypeError, ValueError, OverflowError):
            return None
        if isinstance(convertido, (datetime, date)):
            return convertido.date().isoformat() if isinstance(convertido, datetime) else convertido.isoformat()
        return None
    texto = str(valor).strip()
    for formato in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(texto, formato).date().isoformat()
        except ValueError:
            continue
    return None


def importar_trabajadores_excel(ruta_excel):
    """
    Procesa un archivo Excel para importar o actualizar el catálogo de trabajadores.
    Retorna un diccionario con los resultados del proceso.
    """
    # Cargar el archivo Excel, ignorando fórmulas y leyendo solo datos
    wb = openpyxl.load_workbook(ruta_excel, data_only=True)

    sheet_seleccionada = None
    fila_cabecera = None
    col_map = {}

    # Buscar la cabecera inspeccionando las primeras 20 filas en todas las hojas
    for sheet in wb.worksheets:
        for row_idx, row in enumerate(sheet.iter_rows(max_row=20, values_only=True), start=1):
            row_clean = [limpiar_texto(c) for c in row if c is not None]
            texto_unido = " ".join(row_clean)

            # Si detecta palabras clave que sugieren que es la tabla principal
            if "DNI" in texto_unido or ("PATERNO" in texto_unido and "MATERNO" in texto_unido):
                sheet_seleccionada = sheet
                fila_cabecera = row_idx

                # Mapear los índices de columnas basados en el nombre de la cabecera
                for idx, cell_val in enumerate(row):
                    h = limpiar_texto(cell_val)
                    if not h:
                        continue

                    h_norm = h.lower().replace(" ", "_")

                    if "dni" in h_norm or "documento" in h_norm or "identidad" in h_norm:
                        col_map["dni"] = idx
                    elif "paterno" in h_norm:
                        col_map["paterno"] = idx
                    elif "materno" in h_norm:
                        col_map["materno"] = idx
                    elif "nombre" in h_norm and "paterno" not in h_norm and "materno" not in h_norm:
                        col_map["nombres"] = idx
                    elif "cargo" in h_norm or "puesto" in h_norm:
                        col_map["cargo"] = idx
                    elif "correo" in h_norm or "email" in h_norm:
                        if "institucional" in h_norm:
                            col_map["correo_institucional"] = idx
                        elif "personal" in h_norm or "privado" in h_norm:
                            col_map["correo_personal"] = idx
                    elif "contrato" in h_norm or "regimen" in h_norm:
                        col_map["tipo_contrato"] = idx
                    elif "telefono" in h_norm or "celular" in h_norm or "movil" in h_norm or "numero_institucional" in h_norm:
                        if "personal" in h_norm or "privado" in h_norm:
                            col_map["telefono"] = idx
                        elif "institucional" in h_norm or "numero_institucional" in h_norm:
                            col_map["celular_institucional"] = idx
                        elif "celular" in h_norm:
                            col_map["celular_institucional"] = idx
                        else:
                            col_map["telefono"] = idx
                    elif "nacimiento" in h_norm:
                        col_map["fecha_nacimiento"] = idx
                break
        if sheet_seleccionada:
            break

    # Fallback: si no encontró cabecera clara, asume la hoja activa y la fila 1
    if not sheet_seleccionada or "dni" not in col_map:
        sheet_seleccionada = wb.active
        fila_cabecera = 1

    conn = get_db()
    cursor = conn.cursor()
    total_nuevos = 0
    total_actualizados = 0
    diferencias = []
    celulares_duplicados = []
    dni_vistos = set()
    filas_invalidas = 0

    # Procesar desde la fila siguiente a la cabecera
    for row in sheet_seleccionada.iter_rows(min_row=fila_cabecera + 1, values_only=True):
        if not row:
            continue

        dni_idx = col_map.get("dni")
        if dni_idx is None or dni_idx >= len(row):
            continue

        # Validar el DNI como llave primaria
        dni_val = row[dni_idx]
        dni = normalizar_valor_excel(dni_val)
        if not dni or dni == "0" or dni.upper() == "NONE":
            continue
        if not re.fullmatch(r"[0-9]{8}", dni):
            filas_invalidas += 1
            continue

        # Evitar procesamiento duplicado del mismo DNI en el mismo excel
        if dni in dni_vistos:
            continue
        dni_vistos.add(dni)

        # Función auxiliar para extraer datos usando el mapa de columnas
        def get_val(campo):
            idx = col_map.get(campo)
            if idx is not None and idx < len(row) and row[idx] is not None:
                val_str = normalizar_valor_excel(row[idx])
                return val_str
            return ""

        nombres = get_val("nombres")
        paterno = get_val("paterno")
        materno = get_val("materno")
        cargo = get_val("cargo")
        c_pers = get_val("correo_personal")
        c_inst = get_val("correo_institucional")
        contrato = get_val("tipo_contrato")
        telefono = get_val("telefono")
        celular_institucional = get_val("celular_institucional")

        raw_f_nac = row[col_map.get("fecha_nacimiento")] if col_map.get("fecha_nacimiento") is not None else None
        f_nac = fecha_iso_importada(raw_f_nac)
        if (
            (telefono and not re.fullmatch(r"[0-9]{1,15}", telefono))
            or (celular_institucional and not re.fullmatch(r"[0-9]{1,15}", celular_institucional))
            or f_nac is None
        ):
            filas_invalidas += 1
            continue

        # Buscar si el trabajador ya existe en BD
        existente = cursor.execute("SELECT * FROM trabajadores WHERE dni = ?", (dni,)).fetchone()
        datos_importados = {
            "dni": dni,
            "nombres": nombres,
            "apellido_paterno": paterno,
            "apellido_materno": materno,
            "cargo": cargo,
            "correo_personal": c_pers,
            "correo_institucional": c_inst,
            "tipo_contrato": contrato,
            "telefono": telefono,
            "celular_institucional": celular_institucional,
            "fecha_nacimiento": f_nac,
        }

        if existente is None:
            # Validar que el celular institucional no esté en uso por otro trabajador
            celular_normalizado = normalizar_numero_celular(celular_institucional)
            trabajador_con_celular = cursor.execute("""
                SELECT dni FROM trabajadores
                WHERE celular_institucional_normalizado = ?
            """, (celular_normalizado,)).fetchone() if celular_normalizado else None
            
            if trabajador_con_celular:
                celulares_duplicados.append({
                    "dni": dni,
                    "dni_existente": trabajador_con_celular["dni"],
                    "celular": celular_institucional,
                })
                continue
                
            # Inserción de un nuevo trabajador
            cursor.execute("""
                INSERT INTO trabajadores (dni, nombres, apellido_paterno, apellido_materno, cargo,
                                          correo_personal, correo_institucional, tipo_contrato,
                                          telefono, celular_institucional, fecha_nacimiento, activo,
                                          celular_institucional_normalizado)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 1, ?)
            """, (
                dni,
                nombres,
                paterno,
                materno,
                cargo,
                c_pers,
                c_inst,
                contrato,
                telefono,
                celular_institucional,
                f_nac,
                celular_normalizado,
            ))
            total_nuevos += 1
            continue

        # Si el trabajador existe, preparamos datos combinados (nuevos sobrescriben solo si no están vacíos)
        datos_nuevos = {
            clave: valor if str(valor).strip() else (existente[clave] or "")
            for clave, valor in datos_importados.items()
        }
        
        # Validar duplicados de celular para actualizaciones
        celular_normalizado = normalizar_numero_celular(datos_nuevos["celular_institucional"])
        trabajador_con_celular = cursor.execute("""
            SELECT dni FROM trabajadores
            WHERE celular_institucional_normalizado = ? AND dni != ?
        """, (celular_normalizado, dni)).fetchone() if celular_normalizado else None
        
        if trabajador_con_celular:
            celulares_duplicados.append({
                "dni": dni,
                "dni_existente": trabajador_con_celular["dni"],
                "celular": datos_nuevos["celular_institucional"],
            })
            continue
            
        # Determinar si hay cambios reales para guardar log
        cambios = []
        for clave, valor in datos_nuevos.items():
            if clave == "dni":
                continue
            valor_actual = (existente[clave] or "")
            if str(valor_actual).strip() != str(valor).strip():
                cambios.append(f"{clave}: '{valor_actual}' -> '{valor}'")

        if cambios:
            # Ejecutar la actualización en BD
            cursor.execute("""
                UPDATE trabajadores
                SET nombres = ?, apellido_paterno = ?, apellido_materno = ?, cargo = ?,
                    correo_personal = ?, correo_institucional = ?, tipo_contrato = ?,
                    telefono = ?, celular_institucional = ?, fecha_nacimiento = ?, activo = 1,
                    celular_institucional_normalizado = ?
                WHERE dni = ?
            """, (
            datos_nuevos["nombres"],
            datos_nuevos["apellido_paterno"],
            datos_nuevos["apellido_materno"],
            datos_nuevos["cargo"],
            datos_nuevos["correo_personal"],
            datos_nuevos["correo_institucional"],
            datos_nuevos["tipo_contrato"],
            datos_nuevos["telefono"],
                datos_nuevos["celular_institucional"],
            datos_nuevos["fecha_nacimiento"],
                celular_normalizado,
                dni,
            ))
            diferencias.append({
                "dni": dni,
                "nombre": f"{nombres} {paterno} {materno}".strip(),
                "cambios": cambios,
            })
            total_actualizados += 1

    conn.commit()
    conn.close()

    return {
        "total": total_nuevos + total_actualizados,
        "nuevos": total_nuevos,
        "actualizados": total_actualizados,
        "mismatches": diferencias,
        "celulares_duplicados": celulares_duplicados,
        "filas_invalidas": filas_invalidas,
    }