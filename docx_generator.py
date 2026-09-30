import os
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor

MESES_ES = [
    'ENERO', 'FEBRERO', 'MARZO', 'ABRIL', 'MAYO', 'JUNIO',
    'JULIO', 'AGOSTO', 'SETIEMBRE', 'OCTUBRE', 'NOVIEMBRE', 'DICIEMBRE'
]

NOMBRES_CATEGORIAS_TDR = {
    'HARDWARE': 'SOPORTE Y MANTENIMIENTO DE HARDWARE / BIENES INFORMÁTICOS',
    'SOFTWARE': 'SOPORTE DE SOFTWARE, APLICATIVOS Y SISTEMAS INSTITUCIONALES',
    'EQUIPOS MOVILES': 'GESTIÓN Y ASISTENCIA DE EQUIPOS MÓVILES INSTITUCIONALES',
    'TOKENS': 'CONFIGURACIÓN Y EMISIÓN DE TOKENS / FIRMA DIGITAL',
    'CAPACITACION': 'CAPACITACIONES, VIDEOCONFERENCIAS Y ASISTENCIAS VIRTUALES',
    'TRAMITE UTI': 'TRÁMITES, INFORMES Y GESTIONES TÉCNICAS UTI'
}

def formatear_fecha_texto_es(fecha_str):
    """Convierte YYYY-MM-DD o DD/MM/YYYY a formato '15 DE SETIEMBRE DEL 2026'."""
    if not fecha_str:
        return ""
    try:
        if "-" in fecha_str:
            y, m, d = [int(x) for x in fecha_str.split("-")]
        elif "/" in fecha_str:
            d, m, y = [int(x) for x in fecha_str.split("/")]
        else:
            return fecha_str
        nombre_mes = MESES_ES[m - 1]
        return f"{d} DE {nombre_mes} DEL {y}"
    except Exception:
        return fecha_str

def procesar_parrafo(parrafo, datos_reemplazo):
    if not parrafo.text:
        return

    hay_etiquetas = any(clave in parrafo.text for clave in datos_reemplazo)
    if not hay_etiquetas:
        return

    for clave, valor in datos_reemplazo.items():
        if clave in parrafo.text:
            for run in parrafo.runs:
                if clave in run.text:
                    run.text = run.text.replace(clave, str(valor or ""))

    if any(clave in parrafo.text for clave in datos_reemplazo):
        texto_completo = parrafo.text
        for clave, valor in datos_reemplazo.items():
            texto_completo = texto_completo.replace(clave, str(valor or ""))
        
        if parrafo.runs:
            parrafo.runs[0].text = texto_completo
            for r in parrafo.runs[1:]:
                r.text = ""
        else:
            parrafo.text = texto_completo

def rellenar_plantilla_docx(ruta_plantilla, datos_reemplazo, ruta_salida):
    doc = Document(ruta_plantilla)

    for p in doc.paragraphs:
        procesar_parrafo(p, datos_reemplazo)

    for tabla in doc.tables:
        for fila in tabla.rows:
            for celda in fila.cells:
                for p in celda.paragraphs:
                    procesar_parrafo(p, datos_reemplazo)

    for seccion in doc.sections:
        for p in seccion.header.paragraphs:
            procesar_parrafo(p, datos_reemplazo)
        for p in seccion.footer.paragraphs:
            procesar_parrafo(p, datos_reemplazo)

    os.makedirs(os.path.dirname(ruta_salida), exist_ok=True)
    doc.save(ruta_salida)
    return ruta_salida

def generar_informe_mensual_docx(ruta_plantilla, num_carta, num_informe, fecha_doc_str, num_os, anio, actividades_lista, ruta_salida):
    """
    Genera el informe mensual TDR en DOCX con tipografía Century Gothic,
    reemplazando encabezados correlativos y agregando al final el anexo
    de todas las actividades agrupadas por temática/categoría.
    """
    if not os.path.exists(ruta_plantilla):
        raise FileNotFoundError(f"No se encontró la plantilla en {ruta_plantilla}")

    doc = Document(ruta_plantilla)
    fecha_formateada = formatear_fecha_texto_es(fecha_doc_str)

    str_carta = f"CARTA N° {int(num_carta):03d}-{anio}-VCEM-SI"
    str_informe = f"INFORME N° {int(num_informe):03d}-{anio}-VCEM-SI"
    str_adjunto = f"Informe N° {int(num_informe):03d}-{anio}-UTAPUR-VCEM-SI."
    str_os = f"ORDEN DE SERVICIO N° {num_os}"
    str_fecha_carta = f"FECHA\t\t:\t{fecha_formateada}"
    str_fecha_informe = f"FECHA\t:\t{fecha_formateada}"

    # 1. Reemplazo de correlativos y fechas en el cuerpo base
    for p in doc.paragraphs:
        txt = p.text
        if not txt:
            continue

        if "CARTA N" in txt and "VCEM-SI" in txt:
            p.text = re.sub(r'CARTA\s+N[°ºo]?\s*[\d\w\-]+VCEM-SI', str_carta, txt, flags=re.IGNORECASE)
            for r in p.runs:
                r.font.name = 'Century Gothic'
        elif "INFORME N" in txt and "VCEM-SI" in txt:
            p.text = re.sub(r'INFORME\s+N[°ºo]?\s*[\d\w\-]+VCEM-SI', str_informe, txt, flags=re.IGNORECASE)
            for r in p.runs:
                r.font.name = 'Century Gothic'
        elif "Informe N" in txt and "UTAPUR" in txt:
            p.text = re.sub(r'Informe\s+N[°ºo]?\s*[\d\w\-]+UTAPUR-VCEM-SI\.?', str_adjunto, txt, flags=re.IGNORECASE)
            for r in p.runs:
                r.font.name = 'Century Gothic'
        elif "ORDEN DE SERVICIO" in txt and num_os:
            p.text = re.sub(r'ORDEN\s+DE\s+SERVICIO\s+N[°ºo]?\s*[\d\w\-]+', str_os, txt, flags=re.IGNORECASE)
            for r in p.runs:
                r.font.name = 'Century Gothic'
        elif "FECHA" in txt and ("AGOSTO" in txt or "DEL 2026" in txt or "DEL 202" in txt):
            if "\t\t:" in txt:
                p.text = str_fecha_carta
            elif "\t:" in txt or ":" in txt:
                p.text = str_fecha_informe
            for r in p.runs:
                r.font.name = 'Century Gothic'

    # 2. Anexo de Actividades TI agrupadas por temática (si hay actividades)
    if actividades_lista:
        p_sep = doc.add_paragraph()
        p_sep.paragraph_format.space_before = Pt(24)

        p_tit = doc.add_paragraph()
        r_tit = p_tit.add_run('ANEXO: RELACIÓN DETALLADA DE ACTIVIDADES TI REGISTRADAS EN EL PERIODO')
        r_tit.font.name = 'Century Gothic'
        r_tit.font.size = Pt(11.5)
        r_tit.bold = True
        r_tit.font.color.rgb = RGBColor(7, 89, 133)
        p_tit.paragraph_format.space_after = Pt(4)

        p_sub = doc.add_paragraph()
        r_sub = p_sub.add_run('Actividades de soporte técnico, mantenimiento y gestión registradas en el sistema, agrupadas por temática:')
        r_sub.font.name = 'Century Gothic'
        r_sub.font.size = Pt(9.5)
        r_sub.italic = True
        p_sub.paragraph_format.space_after = Pt(10)

        # Agrupar por categoría
        grupos = {}
        for act in actividades_lista:
            cat = act.get('categoria') or 'OTRAS ACTIVIDADES'
            grupos.setdefault(cat, []).append(act)

        for cat_key, items in grupos.items():
            nombre_cat = NOMBRES_CATEGORIAS_TDR.get(cat_key, cat_key.upper())

            # Título de Categoría / Tema
            p_cat = doc.add_paragraph()
            p_cat.paragraph_format.space_before = Pt(10)
            p_cat.paragraph_format.space_after = Pt(3)
            r_cat = p_cat.add_run(f'■ {nombre_cat} ({len(items)} actividades)')
            r_cat.font.name = 'Century Gothic'
            r_cat.font.size = Pt(10.5)
            r_cat.bold = True
            r_cat.font.color.rgb = RGBColor(15, 23, 42)

            # Lista de actividades del grupo
            for idx, act in enumerate(items, 1):
                fecha_item = act.get('fecha_programada') or '-'
                titulo_item = (act.get('titulo') or '').strip().upper()
                desc_item = (act.get('descripcion') or '').strip()
                sol_item = (act.get('solucion_aplicada') or '').strip()
                area_item = (act.get('area_solicitante') or '').strip()

                # Línea 1: Número, Fecha y Título
                p_item = doc.add_paragraph()
                p_item.paragraph_format.left_indent = Inches(0.25)
                p_item.paragraph_format.space_after = Pt(1)

                r_n = p_item.add_run(f'{idx}. ')
                r_n.font.name = 'Century Gothic'
                r_n.font.size = Pt(9.5)
                r_n.bold = True

                r_f = p_item.add_run(f'[{fecha_item}] ')
                r_f.font.name = 'Century Gothic'
                r_f.font.size = Pt(9)
                r_f.bold = True
                r_f.font.color.rgb = RGBColor(7, 89, 133)

                r_t = p_item.add_run(titulo_item)
                r_t.font.name = 'Century Gothic'
                r_t.font.size = Pt(9.5)
                r_t.bold = True

                # Línea 2: Acción Realizada / Detalle
                texto_detalle = sol_item if sol_item else (desc_item if desc_item else 'Actividad ejecutada y concluida satisfactoriamente.')
                p_det = doc.add_paragraph()
                p_det.paragraph_format.left_indent = Inches(0.45)
                p_det.paragraph_format.space_after = Pt(4)

                r_lbl = p_det.add_run('Acción / Detalle: ')
                r_lbl.font.name = 'Century Gothic'
                r_lbl.font.size = Pt(9)
                r_lbl.italic = True
                r_lbl.font.color.rgb = RGBColor(71, 85, 105)

                r_txt = p_det.add_run(texto_detalle)
                r_txt.font.name = 'Century Gothic'
                r_txt.font.size = Pt(9)

                if area_item and area_item != 'UT Apurímac':
                    r_ar = p_det.add_run(f' (Área: {area_item})')
                    r_ar.font.name = 'Century Gothic'
                    r_ar.font.size = Pt(8.5)
                    r_ar.italic = True

    os.makedirs(os.path.dirname(ruta_salida), exist_ok=True)
    doc.save(ruta_salida)
    return ruta_salida

def generar_reporte_actividades_docx(actividades_lista, periodo_desde, periodo_hasta, categoria_filtro, ruta_salida):
    """
    Genera un documento Word independiente con todas las actividades TI
    agrupadas por categoría temática y ordenadas cronológicamente, con formato
    Century Gothic listo para copiar y pegar en informes de actividades.
    """
    doc = Document()

    # Márgenes de página estándar
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.9)
        s.right_margin = Inches(0.9)

    # Encabezado Institucional
    p_inst = doc.add_paragraph()
    r_inst = p_inst.add_run('PROGRAMA NACIONAL DE ALIMENTACIÓN ESCOLAR PAE - UT APURÍMAC\n')
    r_inst.font.name = 'Century Gothic'
    r_inst.font.size = Pt(9.5)
    r_inst.font.color.rgb = RGBColor(100, 116, 139)
    r_inst.bold = True

    r_main = p_inst.add_run('REPORTE CONSOLIDADO DE ACTIVIDADES TI (SOPORTE INFORMÁTICO)')
    r_main.font.name = 'Century Gothic'
    r_main.font.size = Pt(13)
    r_main.bold = True
    r_main.font.color.rgb = RGBColor(7, 89, 133)
    p_inst.paragraph_format.space_after = Pt(2)

    # Subtítulo y Metadatos
    p_sub = doc.add_paragraph()
    sub_txt = 'Organizado por Categorías Temáticas y Orden Cronológico para Informe Mensual'
    if periodo_desde or periodo_hasta:
        sub_txt += f' | Periodo: {periodo_desde or "Inicio"} al {periodo_hasta or "Actualidad"}'
    sub_txt += f' | Total: {len(actividades_lista)} actividades registradas'
    r_sub = p_sub.add_run(sub_txt)
    r_sub.font.name = 'Century Gothic'
    r_sub.font.size = Pt(9.5)
    r_sub.italic = True
    p_sub.paragraph_format.space_after = Pt(14)

    # Agrupación por Categoría
    grupos = {}
    for act in actividades_lista:
        cat = act.get('categoria') or 'OTRAS ACTIVIDADES'
        grupos.setdefault(cat, []).append(act)

    for cat_key, items in grupos.items():
        nombre_cat = NOMBRES_CATEGORIAS_TDR.get(cat_key, cat_key.upper())

        # Título de Categoría
        p_cat = doc.add_paragraph()
        p_cat.paragraph_format.space_before = Pt(12)
        p_cat.paragraph_format.space_after = Pt(4)
        r_cat = p_cat.add_run(f'■ {nombre_cat} ({len(items)} actividades)')
        r_cat.font.name = 'Century Gothic'
        r_cat.font.size = Pt(11)
        r_cat.bold = True
        r_cat.font.color.rgb = RGBColor(15, 23, 42)

        # Actividades ordenadas por fecha
        items_ordenados = sorted(items, key=lambda x: (x.get('fecha_programada') or '', x.get('hora_inicio') or ''))

        for idx, act in enumerate(items_ordenados, 1):
            fecha_item = act.get('fecha_programada') or '-'
            titulo_item = (act.get('titulo') or '').strip().upper()
            desc_item = (act.get('descripcion') or '').strip()
            sol_item = (act.get('solucion_aplicada') or '').strip()
            area_item = (act.get('area_solicitante') or '').strip()
            prio_item = act.get('prioridad') or ''
            est_item = act.get('estado') or ''
            h_ini = act.get('hora_inicio') or ''
            h_fin = act.get('hora_fin') or ''

            # Línea 1: Número, Fecha y Título
            p_item = doc.add_paragraph()
            p_item.paragraph_format.left_indent = Inches(0.2)
            p_item.paragraph_format.space_after = Pt(1)

            r_n = p_item.add_run(f'{idx}. ')
            r_n.font.name = 'Century Gothic'
            r_n.font.size = Pt(9.5)
            r_n.bold = True

            r_f = p_item.add_run(f'[{fecha_item}] ')
            r_f.font.name = 'Century Gothic'
            r_f.font.size = Pt(9)
            r_f.bold = True
            r_f.font.color.rgb = RGBColor(7, 89, 133)

            r_t = p_item.add_run(titulo_item)
            r_t.font.name = 'Century Gothic'
            r_t.font.size = Pt(9.5)
            r_t.bold = True

            # Línea 2: Metadatos y Solución
            p_det = doc.add_paragraph()
            p_det.paragraph_format.left_indent = Inches(0.4)
            p_det.paragraph_format.space_after = Pt(4)

            meta = []
            if area_item: meta.append(f'Área: {area_item}')
            if prio_item: meta.append(f'Prioridad: {prio_item}')
            if est_item: meta.append(f'Estado: {est_item}')
            if h_ini or h_fin: meta.append(f'Horario: {h_ini}-{h_fin}')

            if meta:
                r_m = p_det.add_run(' • ' + ' | '.join(meta) + '\n')
                r_m.font.name = 'Century Gothic'
                r_m.font.size = Pt(8.5)
                r_m.italic = True
                r_m.font.color.rgb = RGBColor(100, 116, 139)

            texto_detalle = sol_item if sol_item else (desc_item if desc_item else 'Actividad ejecutada y concluida satisfactoriamente.')
            r_lbl = p_det.add_run(' • Acción Realizada / Detalle: ')
            r_lbl.font.name = 'Century Gothic'
            r_lbl.font.size = Pt(9)
            r_lbl.bold = True

            r_txt = p_det.add_run(texto_detalle)
            r_txt.font.name = 'Century Gothic'
            r_txt.font.size = Pt(9)

    os.makedirs(os.path.dirname(ruta_salida), exist_ok=True)
    doc.save(ruta_salida)
    return ruta_salida