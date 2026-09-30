import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import ColorScaleRule
from datetime import datetime, date
from collections import defaultdict
from io import BytesIO

def aplicar_destinatarios_ocultos(resultado, ocultos):
    """Devuelve una vista del corte sin destinatarios ocultos, sin mutar el original."""
    ocultos = set(ocultos or [])
    visibles = [p for p in resultado.get("destinatarios", []) if p not in ocultos]
    resultado_visible = dict(resultado)
    matriz = resultado.get("matriz", {})
    totales = resultado.get("totales_destinatario", {})
    resultado_visible["destinatarios"] = visibles
    resultado_visible["matriz"] = {p: dict(matriz.get(p, {})) for p in visibles}
    resultado_visible["totales_destinatario"] = {p: totales.get(p, 0) for p in visibles}
    resultado_visible["totales_fecha"] = {
        fecha: sum(resultado_visible["matriz"].get(p, {}).get(fecha, 0) for p in visibles)
        for fecha in resultado.get("fechas", [])
    }
    resultado_visible["gran_total"] = sum(resultado_visible["totales_destinatario"].values())
    resultado_visible["total_trabajadores"] = len(visibles)
    resultado_visible["total_min"] = min(resultado_visible["totales_destinatario"].values(), default=0)
    resultado_visible["total_max"] = max(resultado_visible["totales_destinatario"].values(), default=0)
    ranking = sorted(resultado_visible["totales_destinatario"].items(), key=lambda item: (item[1], item[0]), reverse=True)
    resultado_visible["ranking_descendente"] = ranking
    resultado_visible["top_10"] = [
        {"pos": pos, "nombre": nombre, "cant": cant,
         "pct": round((cant / resultado_visible["gran_total"] * 100), 1) if resultado_visible["gran_total"] else 0}
        for pos, (nombre, cant) in enumerate(ranking[:10], start=1)
    ]
    resultado_visible["suma_top10"] = sum(item["cant"] for item in resultado_visible["top_10"])
    resultado_visible["pct_top10"] = round((resultado_visible["suma_top10"] / resultado_visible["gran_total"] * 100), 1) if resultado_visible["gran_total"] else 0
    documentos_100_dias = [
        item for item in resultado.get("documentos_100_dias", [])
        if item["nombre"] not in ocultos
    ]
    resultado_visible["documentos_100_dias"] = documentos_100_dias
    resultado_visible["total_documentos_100_dias"] = sum(item["cant"] for item in documentos_100_dias)
    resultado_visible["total_trabajadores_100_dias"] = len(documentos_100_dias)
    return resultado_visible

def convertir_a_fecha_dma(valor_raw):
    """
    Convierte cualquier valor textual del campo FECHA EMI.
    a formato de fecha DMA (DD/MM/YYYY) y genera clave cronológica YYYYMMDD.
    """
    if not valor_raw:
        return None, None

    if isinstance(valor_raw, (datetime, date)):
        return valor_raw.strftime("%d/%m/%Y"), valor_raw.strftime("%Y%m%d")

    texto = str(valor_raw).strip()
    if not texto or texto.upper() in ("NONE", "NULL", ""):
        return None, None

    partes_espacio = texto.split()
    if not partes_espacio:
        return None, None
    parte_fecha = partes_espacio[0].strip()

    if "/" in parte_fecha:
        segmentos = parte_fecha.split("/")
        if len(segmentos) == 3:
            s_primero, s_medio, s_ultimo = [str(x).strip() for x in segmentos]
            try:
                if len(s_primero) == 4:
                    yyyy = int(s_primero)
                    mm = int(s_medio)
                    dd = int(s_ultimo)
                else:
                    dd = int(s_primero)
                    mm = int(s_medio)
                    yyyy = int(s_ultimo)
                dt = datetime(yyyy, mm, dd)
                return dt.strftime("%d/%m/%Y"), dt.strftime("%Y%m%d")
            except Exception:
                pass

    if "-" in parte_fecha:
        segmentos = parte_fecha.split("-")
        if len(segmentos) == 3:
            s_primero, s_medio, s_ultimo = [str(x).strip() for x in segmentos]
            try:
                if len(s_primero) == 4:
                    yyyy = int(s_primero)
                    mm = int(s_medio)
                    dd = int(s_ultimo)
                else:
                    dd = int(s_primero)
                    mm = int(s_medio)
                    yyyy = int(s_ultimo)
                dt = datetime(yyyy, mm, dd)
                return dt.strftime("%d/%m/%Y"), dt.strftime("%Y%m%d")
            except Exception:
                pass

    return parte_fecha, parte_fecha

def calcular_antiguedad_dias(fecha_dma_str, dt_referencia):
    try:
        dd, mm, yyyy = [int(x) for x in fecha_dma_str.split("/")]
        dt_doc = datetime(yyyy, mm, dd)
        diff = (dt_referencia - dt_doc).days
        return max(0, diff)
    except Exception:
        return 0

def procesar_sgd_excel(ruta_archivo, tipo_estado="NO LEIDOS", fecha_corte_str=None):
    """
    Construye la Tabla Dinámica SGD para NO LEÍDOS o RECIBIDOS:
    - FILAS: DESTINATARIO
    - COLUMNAS: FECHA EMI. (convertida a DMA)
    - VALORES: Cuenta de ESTADO
    """
    if not os.path.exists(ruta_archivo):
        return None

    wb = openpyxl.load_workbook(ruta_archivo, data_only=True)
    ws = wb.active

    col_destinatario = None
    col_fecha_emi = None
    header_row = None

    for r in range(1, min(30, ws.max_row + 1)):
        row_cells = [str(ws.cell(row=r, column=c).value or "").strip().upper() for c in range(1, ws.max_column + 1)]
        destinatario_candidato = None
        fecha_emision_candidata = None
        fecha_generica_candidata = None

        for idx, val in enumerate(row_cells, start=1):
            if destinatario_candidato is None and any(k in val for k in ["DESTINATARIO", "DESTINO", "USUARIO DESTINO", "RECEPTOR", "NOMBRE DEL DESTINATARIO", "PERSONA", "NOMBRE"]):
                destinatario_candidato = idx
            if fecha_emision_candidata is None and any(k in val for k in ["FECHA EMI", "F. EMI", "F.EMI", "FECHA DE EMI", "FECHA_EMI", "FECHA EMIS", "F.EMIS"]):
                fecha_emision_candidata = idx
            elif fecha_generica_candidata is None and "FECHA" in val:
                fecha_generica_candidata = idx

        fecha_candidata = fecha_emision_candidata or fecha_generica_candidata
        if destinatario_candidato and fecha_candidata:
            col_destinatario = destinatario_candidato
            col_fecha_emi = fecha_candidata
            header_row = r
            break

    header_row = header_row or 1
    col_destinatario = col_destinatario or 3
    col_fecha_emi = col_fecha_emi or 4

    tabla_dinamica = defaultdict(lambda: defaultdict(int))
    mapa_orden_fechas = {}
    totales_por_destinatario = defaultdict(int)

    for r in range(header_row + 1, ws.max_row + 1):
        dest_val = ws.cell(row=r, column=col_destinatario).value
        fecha_val = ws.cell(row=r, column=col_fecha_emi).value

        if not dest_val or str(dest_val).strip() == "":
            continue

        destinatario = " ".join(str(dest_val).strip().upper().split())
        
        if any(destinatario.startswith(k) for k in ["TOTAL", "SUBTOTAL", "ETIQUETAS"]):
            continue

        fecha_dma, fecha_sort = convertir_a_fecha_dma(fecha_val)
        if not fecha_dma:
            fecha_dma = "SIN FECHA"
            fecha_sort = "00000000"

        tabla_dinamica[destinatario][fecha_dma] += 1
        totales_por_destinatario[destinatario] += 1
        mapa_orden_fechas[fecha_dma] = fecha_sort

    fechas_columnas = sorted(mapa_orden_fechas.keys(), key=lambda f: mapa_orden_fechas.get(f, "00000000"))
    destinatarios_filas = sorted(totales_por_destinatario.keys(), key=lambda d: (totales_por_destinatario.get(d, 0), d))

    totales_por_fecha = {}
    for f in fechas_columnas:
        totales_por_fecha[f] = sum(tabla_dinamica[d].get(f, 0) for d in destinatarios_filas)

    gran_total = sum(totales_por_destinatario.values())
    total_trabajadores = len(destinatarios_filas)

    ranking_descendente = sorted(totales_por_destinatario.items(), key=lambda x: (x[1], x[0]), reverse=True)
    top_10 = ranking_descendente[:10]

    top_10_detalle = []
    for pos, (nombre, cant) in enumerate(top_10, start=1):
        pct = round((cant / gran_total * 100), 1) if gran_total > 0 else 0
        top_10_detalle.append({
            "pos": pos,
            "nombre": nombre,
            "cant": cant,
            "pct": pct
        })

    suma_top10 = sum(item["cant"] for item in top_10_detalle)
    pct_top10 = round((suma_top10 / gran_total * 100), 1) if gran_total > 0 else 0

    dt_corte = datetime.now()
    if fecha_corte_str:
        try:
            f_parte = fecha_corte_str.split()[0]
            if "/" in f_parte:
                d_c, m_c, y_c = [int(x) for x in f_parte.split("/")]
                dt_corte = datetime(y_c, m_c, d_c)
        except Exception:
            dt_corte = datetime.now()

    rangos_antiguedad = {
        "0 a 2 días (Recientes)": {"cant": 0, "pct": 0, "color": "success"},
        "3 a 5 días (Atención Normal)": {"cant": 0, "pct": 0, "color": "info"},
        "6 a 10 días (Alerta Media)": {"cant": 0, "pct": 0, "color": "warning"},
        "11 a 20 días (Alerta Alta)": {"cant": 0, "pct": 0, "color": "danger"},
        "Más de 20 días (Críticos / Muy Antiguos)": {"cant": 0, "pct": 0, "color": "dark"}
    }

    for f_dma, cant in totales_por_fecha.items():
        if f_dma == "SIN FECHA":
            rangos_antiguedad["Más de 20 días (Críticos / Muy Antiguos)"]["cant"] += cant
            continue
        dias = calcular_antiguedad_dias(f_dma, dt_corte)
        if dias <= 2:
            rangos_antiguedad["0 a 2 días (Recientes)"]["cant"] += cant
        elif dias <= 5:
            rangos_antiguedad["3 a 5 días (Atención Normal)"]["cant"] += cant
        elif dias <= 10:
            rangos_antiguedad["6 a 10 días (Alerta Media)"]["cant"] += cant
        elif dias <= 20:
            rangos_antiguedad["11 a 20 días (Alerta Alta)"]["cant"] += cant
        else:
            rangos_antiguedad["Más de 20 días (Críticos / Muy Antiguos)"]["cant"] += cant

    for r_k, r_v in rangos_antiguedad.items():
        r_v["pct"] = round((r_v["cant"] / gran_total * 100), 1) if gran_total > 0 else 0

    documentos_100_dias = []
    documentos_sin_fecha = 0
    if str(tipo_estado).strip().upper().replace(" ", "_") in {"NO_LEIDOS", "RECIBIDOS"}:
        for destinatario, documentos_por_fecha in tabla_dinamica.items():
            cantidad_vencida = 0
            fecha_mas_antigua = None
            antiguedad_maxima = 0
            for fecha, cantidad in documentos_por_fecha.items():
                try:
                    fecha_documento = datetime.strptime(fecha, "%d/%m/%Y")
                except (TypeError, ValueError):
                    documentos_sin_fecha += cantidad
                    continue
                antiguedad = (dt_corte - fecha_documento).days
                if antiguedad >= 100:
                    cantidad_vencida += cantidad
                    if antiguedad > antiguedad_maxima:
                        antiguedad_maxima = antiguedad
                        fecha_mas_antigua = fecha
            if cantidad_vencida:
                documentos_100_dias.append({
                    "nombre": destinatario,
                    "cant": cantidad_vencida,
                    "fecha_mas_antigua": fecha_mas_antigua,
                    "dias_max": antiguedad_maxima
                })
    documentos_100_dias.sort(key=lambda item: (-item["dias_max"], -item["cant"], item["nombre"]))

    fechas_detalle = []
    for f in fechas_columnas:
        cant_f = totales_por_fecha.get(f, 0)
        pct_f = round((cant_f / gran_total * 100), 1) if gran_total > 0 else 0
        fechas_detalle.append({
            "fecha": f,
            "cant": cant_f,
            "pct": pct_f
        })

    criticos_count = sum(1 for _, c in ranking_descendente if c > 10)
    medios_count = sum(1 for _, c in ranking_descendente if 4 <= c <= 10)
    leves_count = sum(1 for _, c in ranking_descendente if c <= 3)
    promedio = round(gran_total / total_trabajadores, 1) if total_trabajadores > 0 else 0

    return {
        "tipo_estado": tipo_estado,
        "destinatarios": destinatarios_filas,
        "fechas": fechas_columnas,
        "matriz": {d: dict(tabla_dinamica[d]) for d in destinatarios_filas},
        "totales_destinatario": dict(totales_por_destinatario),
        "totales_fecha": totales_por_fecha,
        "gran_total": gran_total,
        "total_trabajadores": total_trabajadores,
        "promedio": promedio,
        "ranking_descendente": ranking_descendente,
        "top_10": top_10_detalle,
        "suma_top10": suma_top10,
        "pct_top10": pct_top10,
        "rangos_antiguedad": rangos_antiguedad,
        "documentos_100_dias": documentos_100_dias,
        "total_documentos_100_dias": sum(item["cant"] for item in documentos_100_dias),
        "total_trabajadores_100_dias": len(documentos_100_dias),
        "documentos_sin_fecha": documentos_sin_fecha,
        "fechas_detalle": fechas_detalle,
        "criticos_count": criticos_count,
        "medios_count": medios_count,
        "leves_count": leves_count
    }

def generar_excel_sgd(resultado, tipo_estado, fecha_corte_str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = tipo_estado.replace(" ", "_")

    ws.merge_cells("A1:Z1")
    title_cell = ws["A1"]
    title_cell.value = f'REPORTE DE DOCUMENTOS EN ESTADO "{tipo_estado}" EN EL SGD AL CORTE {fecha_corte_str}'
    title_cell.font = Font(name="Calibri", size=13, bold=True, color="000000")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions.height = 32

    headers = ["DESTINATARIO"] + resultado["fechas"] + ["Total general"]
    ws.append([])
    ws.append(headers)

    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    header_font = Font(name="Calibri", size=9, bold=True)
    thin_border = Border(
        left=Side(style='thin', color='A6A6A6'), right=Side(style='thin', color='A6A6A6'),
        top=Side(style='thin', color='A6A6A6'), bottom=Side(style='thin', color='A6A6A6')
    )

    ws.row_dimensions.height = 65

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=3, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border
        if col_idx == 1:
            cell.alignment = Alignment(horizontal="left", vertical="center")
        elif col_idx == len(headers):
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="center", vertical="center", text_rotation=90)

    cell_fill_pink = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")

    for r_idx, persona in enumerate(resultado["destinatarios"], start=4):
        total_p = resultado["totales_destinatario"].get(persona, 0)
        row_vals = [persona]
        for f in resultado["fechas"]:
            val = resultado["matriz"].get(persona, {}).get(f, None)
            row_vals.append(val if val else None)
        row_vals.append(total_p)

        ws.append(row_vals)
        ws.row_dimensions[r_idx].height = 20

        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=r_idx, column=col_idx)
            c.border = thin_border
            if col_idx == 1:
                c.alignment = Alignment(horizontal="left", vertical="center")
                c.font = Font(name="Calibri", size=9)
            elif col_idx == len(headers):
                c.alignment = Alignment(horizontal="center", vertical="center")
                c.font = Font(name="Calibri", size=9, bold=True)
            else:
                c.alignment = Alignment(horizontal="center", vertical="center")
                c.font = Font(name="Calibri", size=9)
                if c.value:
                    c.fill = cell_fill_pink

    row_foot = len(resultado["destinatarios"]) + 4
    foot_vals = ["Total general"]
    for f in resultado["fechas"]:
        foot_vals.append(resultado["totales_fecha"].get(f, 0))
    foot_vals.append(resultado["gran_total"])

    ws.append(foot_vals)
    ws.row_dimensions[row_foot].height = 24

    foot_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    for col_idx in range(1, len(headers) + 1):
        c = ws.cell(row=row_foot, column=col_idx)
        c.fill = foot_fill
        c.font = Font(name="Calibri", size=9, bold=True)
        c.border = thin_border
        c.alignment = Alignment(horizontal="center" if col_idx > 1 else "left", vertical="center")

    ws.column_dimensions["A"].width = 38
    for col_idx in range(2, len(headers)):
        col_let = openpyxl.utils.get_column_letter(col_idx)
        ws.column_dimensions[col_let].width = 6
    total_col_letter = openpyxl.utils.get_column_letter(len(headers))
    ws.column_dimensions[total_col_letter].width = 12
    if len(resultado["destinatarios"]):
        ws.conditional_formatting.add(
            f"{total_col_letter}4:{total_col_letter}{row_foot - 1}",
            ColorScaleRule(start_type="min", start_color="C6EFCE",
                           mid_type="percentile", mid_value=50, mid_color="FFF2CC",
                           end_type="max", end_color="F8CBAD")
        )

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer