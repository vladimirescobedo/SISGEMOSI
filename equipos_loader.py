import openpyxl
import re
from database import get_db

def limpiar_texto(txt):
    if txt is None:
        return ""
    txt = str(txt).strip().upper()
    txt = txt.replace("Á", "A").replace("É", "E").replace("Í", "I").replace("Ó", "O").replace("Ú", "U")
    txt = re.sub(r'[\r\n\t]+', ' ', txt)
    txt = re.sub(r'\s+', ' ', txt)
    return txt.strip()

def detectar_categoria(descripcion, marca="", modelo="", observaciones="", nombre_hoja=""):
    hoja_upper = limpiar_texto(nombre_hoja)
    if "MOVIL" in hoja_upper or "CELULAR" in hoja_upper:
        return "MOVILES"
    if "TOKEN" in hoja_upper or "FIRMA" in hoja_upper:
        return "TOKENS"
    if "INFORMATIC" in hoja_upper or "COMPUTO" in hoja_upper:
        return "INFORMATICOS"
    if "MOBILIARIO" in hoja_upper or "MUEBLE" in hoja_upper:
        return "MOBILIARIO"

    texto = limpiar_texto(f"{descripcion} {marca} {modelo} {observaciones}")

    # 1. MOVILES (Celulares, Tablets, Módems móviles)
    patrones_moviles = [
        "CELULAR", "MOVIL", "SMARTPHONE", "TABLET", "TABLETA", "CHIP", "LINEA",
        "TELEFONO CELULAR", "TELEFONO MOVIL", "TELEFONIA", "EQUIPO CELULAR", "MODEM INALAMBRICO"
    ]
    if any(p in texto for p in patrones_moviles):
        return "MOVILES"

    # 2. TOKENS Y SEGURIDAD DIGITAL
    patrones_tokens = [
        "TOKEN", "CRIPTOGRAFICO", "FIRMA DIGITAL", "LLAVE DIGITAL", "CERTIFICADO DIGITAL",
        "DISPOSITIVO CRIPTOGRAFICO", "DISPOSITIVO DE SEGURIDAD", "EPASS", "SAFENET"
    ]
    if any(p in texto for p in patrones_tokens):
        return "TOKENS"

    # 3. MOBILIARIO Y ENSERES DE OFICINA (Sillas, Mesas, Armarios, etc.)
    patrones_mobiliario = [
        "SILLA", "SILLON", "MESA", "ESCRITORIO", "ARMARIO", "ESTANTE", "ARCHIVADOR",
        "CREDENZA", "GAVETA", "VITRINA", "PIZARRA", "PAPELERA", "PERCHERO", "ASIENTO",
        "TABURETE", "BANCA", "BAUL", "MUEBLE", "MODULO", "ESTANTERIA", "CASILLERO",
        "LOCKER", "VENTILADOR", "EXTINTOR", "DISPENSADOR", "HERVIDOR", "MICROONDAS",
        "REFRIGERADOR", "FRIGOBAR", "CALCULADORA", "ENGRAPADOR", "PERFORADOR", "GUILLOTINA",
        "PODIO", "ATRIL", "ESCALERA", "CUADRO", "RELOJ DE PARED", "PORTABANDERA"
    ]
    if any(p in texto for p in patrones_mobiliario):
        return "MOBILIARIO"

    # 4. EQUIPOS INFORMÁTICOS / AUDIO / VIDEO
    patrones_informaticos = [
        "COMPUTADORA", "LAPTOP", "PORTATIL", "NOTEBOOK", "PC", "CPU", "SERVIDOR",
        "MONITOR", "PANTALLA", "CAMARA", "WEBCAM", "TRIPODE", "PROYECTOR", "ECRAN",
        "IMPRESORA", "MULTIFUNCIONAL", "ESCANER", "SCANNER", "DISCO DURO", "MEMORIA USB",
        "PENDRIVE", "ESTABILIZADOR", "UPS", "SWITCH", "ROUTER", "ACCESS POINT", "ACCESSPOINT",
        "RACK", "TECLADO", "MOUSE", "AURICULAR", "HEADSET", "MICROFONO", "PARLANTE", "ALTAVOZ",
        "LECTOR", "BIOMETRICO", "TRANSCEPTOR", "ADAPTADOR", "TRANSFORMADOR", "ANTENA",
        "EQUIPO INFORMATICO", "EQUIPO DE COMPUTO"
    ]
    if any(p in texto for p in patrones_informaticos):
        return "INFORMATICOS"

    # Por defecto, si no es de cómputo, se separa en Mobiliario/Otros
    return "MOBILIARIO"

def importar_equipos_excel(ruta_excel, categoria_forzada=None):
    wb = openpyxl.load_workbook(ruta_excel, data_only=True)
    conn = get_db()
    cursor = conn.cursor()
    
    total_cargados = {"MOVILES": 0, "TOKENS": 0, "INFORMATICOS": 0, "MOBILIARIO": 0}

    for sheet in wb.worksheets:
        fila_cabecera = None
        col_map = {}

        # 1. Detectar cabeceras
        for row_idx, row in enumerate(sheet.iter_rows(max_row=20, values_only=True), start=1):
            row_clean = [limpiar_texto(c) for c in row if c is not None]
            texto_unido = " ".join(row_clean)
            
            if "MARGESI" in texto_unido or "DESCRIPCION" in texto_unido or "SERIE" in texto_unido:
                fila_cabecera = row_idx
                for idx, cell_val in enumerate(row):
                    h = limpiar_texto(cell_val)
                    if not h: continue
                    
                    if "MARGESI" in h or "CODIGO" in h: col_map["codigo_margesi"] = idx
                    elif "DESCRIPCION" in h or "DESCRIPCIÓN" in h: col_map["descripcion"] = idx
                    elif "MARCA" in h: col_map["marca"] = idx
                    elif "MODELO" in h: col_map["modelo"] = idx
                    elif "SERIE" in h or "MEDIDAS" in h: col_map["serie_medidas"] = idx
                    elif "CONSERVAC" in h or "SIGA" in h: col_map["estado_conservacion"] = idx
                    elif "COLOR" in h: col_map["color"] = idx
                    elif "SEDE" in h or "LOCAL" in h: col_map["sede"] = idx
                    elif "INICIAL" in h: col_map["valor_inicial"] = idx
                    elif "NETO" in h: col_map["valor_neto"] = idx
                    elif "OBSERVACION" in h: col_map["observaciones"] = idx
                break

        if not fila_cabecera or "codigo_margesi" not in col_map:
            continue

        # 2. Leer filas y clasificar con precisión
        for row in sheet.iter_rows(min_row=fila_cabecera + 1, values_only=True):
            if not row:
                continue

            cod_idx = col_map.get("codigo_margesi")
            if cod_idx is None or cod_idx >= len(row):
                continue

            cod_val = row[cod_idx]
            if cod_val is None or str(cod_val).strip() == "" or str(cod_val).strip().upper() == "NONE":
                continue

            if isinstance(cod_val, float):
                codigo_margesi = str(int(cod_val)).strip()
            else:
                codigo_margesi = str(cod_val).strip()
                if "." in codigo_margesi:
                    codigo_margesi = codigo_margesi.split(".")[0].strip()

            def get_val(campo):
                idx = col_map.get(campo)
                if idx is not None and idx < len(row) and row[idx] is not None:
                    v = str(row[idx]).strip()
                    return "" if v.upper() == "NONE" else v
                return ""

            descripcion = get_val("descripcion")
            marca = get_val("marca")
            modelo = get_val("modelo")
            serie_medidas = get_val("serie_medidas")
            estado_conservacion = get_val("estado_conservacion")
            color = get_val("color")
            sede = get_val("sede")
            valor_inicial = get_val("valor_inicial")
            valor_neto = get_val("valor_neto")
            observaciones = get_val("observaciones")

            cat = categoria_forzada if categoria_forzada else detectar_categoria(descripcion, marca, modelo, observaciones, sheet.title)

            # Inserción / Actualización
            cursor.execute("SELECT id FROM equipos WHERE codigo_margesi = ? OR codigo_interno = ?", (codigo_margesi, codigo_margesi))
            existe = cursor.fetchone()

            if existe:
                cursor.execute("""
                    UPDATE equipos 
                    SET codigo_margesi = ?, codigo_interno = ?, categoria = ?, tipo_equipo = ?, descripcion = ?, 
                        marca = ?, modelo = ?, serie_medidas = ?, numero_serie = ?,
                        estado_conservacion = ?, color = ?, sede = ?, valor_inicial = ?, valor_neto = ?, observaciones = ?
                    WHERE id = ?
                """, (codigo_margesi, codigo_margesi, cat, descripcion, descripcion, 
                      marca, modelo, serie_medidas, serie_medidas,
                      estado_conservacion, color, sede, valor_inicial, valor_neto, observaciones, existe[0]))
            else:
                cursor.execute("""
                    INSERT INTO equipos (codigo_margesi, codigo_interno, categoria, tipo_equipo, descripcion, 
                                         marca, modelo, serie_medidas, numero_serie,
                                         estado_conservacion, color, sede, valor_inicial, valor_neto, observaciones, estado)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'EN ALMACEN')
                """, (codigo_margesi, codigo_margesi, cat, descripcion, descripcion, 
                      marca, modelo, serie_medidas, serie_medidas,
                      estado_conservacion, color, sede, valor_inicial, valor_neto, observaciones))

            total_cargados[cat] = total_cargados.get(cat, 0) + 1

    conn.commit()
    conn.close()
    return total_cargados